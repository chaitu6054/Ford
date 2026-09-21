"""Lecture et nettoyage du classeur de base (valeurs en cache, couleurs de ligne, liste FSA)."""

from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook

from . import config as C


def nettoyer_nom_colonne(valeur: object) -> str:
    """Nettoyer nom colonne."""
    try:
        if valeur is None:
            return ""
        return " ".join(str(valeur).replace("\n", " ").split()).strip()
    except Exception:
        raise


def trouver_colonne(df: pd.DataFrame, *noms: str) -> str | None:
    """Trouver colonne."""
    try:
        corr = {
            nettoyer_nom_colonne(c).casefold(): c for c in df.columns if nettoyer_nom_colonne(c)
        }
        for nom in noms:
            c = corr.get(nettoyer_nom_colonne(nom).casefold())
            if c is not None:
                return c
        return None
    except Exception:
        raise


def resoudre_colonnes(df: pd.DataFrame, table: dict | None = None) -> dict[str, str | None]:
    """Resoudre colonnes."""
    try:
        return {k: trouver_colonne(df, *v) for k, v in (table or C.COLONNES).items()}
    except Exception:
        raise


def serie_numerique(s: pd.Series) -> pd.Series:
    """Nettoie montants texte : symboles et codes monétaires, espaces unicode
    (dont l'insécable étroite \\u202f des formats français), virgule décimale,
    négatifs comptables entre parenthèses, tirets vides."""
    try:
        if pd.api.types.is_numeric_dtype(s):
            return pd.to_numeric(s, errors="coerce").astype("float64")
        t = (
            s.astype("string")
            .str.replace("€", "", regex=False)
            .str.replace("$", "", regex=False)
            .str.replace("£", "", regex=False)
            .str.replace("%", "", regex=False)
            .str.replace(r"(?i)\b(usd|eur|gbp|cad|chf)\b", "", regex=True)
            .str.replace(r"[\s\u00a0\u202f\u2009\u2007\u3000]", "", regex=True)
            .str.replace(r"^[-–—]$", "", regex=True)
        )
        # "(1 234,56)" (négatif comptable) -> "-1234,56"
        paren = t.str.fullmatch(r"\(.+\)", na=False)
        t = t.where(~paren, "-" + t.str.slice(1, -1))
        t = t.str.replace("+", "", regex=False)
        # Les deux séparateurs présents : le plus à droite est la décimale
        # ("2,026.00" -> "2026.00" ; "1.234,56" -> "1234.56").
        virgule = t.str.contains(",", regex=False, na=False)
        point = t.str.contains(".", regex=False, na=False)
        deux = virgule & point
        virgule_decimale = deux & (t.str.rfind(",") > t.str.rfind("."))
        t = t.where(
            ~virgule_decimale,
            t.str.replace(".", "", regex=False).str.replace(",", ".", regex=False),
        )
        t = t.where(~(deux & ~virgule_decimale), t.str.replace(",", "", regex=False))
        # Que des virgules : une seule -> décimale ("16317,49") ;
        # plusieurs -> milliers ("1,234,567").
        seule_virgule = virgule & ~point
        multi = seule_virgule & (t.str.count(",") > 1)
        t = t.where(~multi, t.str.replace(",", "", regex=False))
        t = t.where(
            ~(seule_virgule & ~multi), t.str.replace(r"(?<=\d),(?=\d)", ".", regex=True)
        )
        return pd.to_numeric(t, errors="coerce").astype("float64")
    except Exception:
        raise


def serie_date(s: pd.Series) -> pd.Series:
    """Datetimes, numéros de série Excel, chaînes. Dates < 1950 => NaT."""
    try:
        num = pd.to_numeric(s, errors="coerce")
        serial = num.where((num > 20000) & (num < 80000))
        depuis_serial = pd.to_datetime(serial, unit="D", origin="1899-12-30", errors="coerce")
        # les cellules numériques hors plage série ne sont jamais des dates (évite l'epoch 1970)
        non_num = s.where(num.isna())
        est_dt = pd.to_datetime(non_num, errors="coerce")
        out = depuis_serial.fillna(est_dt)
        return out.where(out >= pd.Timestamp(C.DATE_MIN_VALIDE))
    except Exception:
        raise


def _rgb(cell) -> str | None:
    try:
        f = cell.fill
        if f is None or f.fill_type != "solid" or f.fgColor is None:
            return None
        rgb = f.fgColor.rgb
        return str(rgb)[-6:].upper() if rgb and rgb != "00000000" else None
    except Exception:
        raise


def lire_code_couleur(chemin: str) -> dict[str, str]:
    """Onglet 'code couleur' : couleur de fond (col A) -> libellé (col B), casefold."""
    try:
        wb = load_workbook(chemin, data_only=True)
        if C.FEUILLE_CODE_COULEUR not in wb.sheetnames:
            return {}
        legende = {}
        for row in wb[C.FEUILLE_CODE_COULEUR].iter_rows(min_row=2, max_col=2):
            rgb, lib = _rgb(row[0]), row[1].value
            if rgb and lib:
                legende[rgb] = nettoyer_nom_colonne(lib).casefold()
        return legende
    except Exception:
        raise


def lire_feuille(chemin: str, feuille: str, legende: dict[str, str] | None = None) -> pd.DataFrame:
    """Feuille avec en-tête ligne 1 ; ajoute `_couleur` (couleur de fond, colonne A)."""
    try:
        wb = load_workbook(chemin, data_only=True)
        ws = wb[feuille]
        lignes = list(ws.iter_rows(min_row=1))
        if not lignes:
            return pd.DataFrame()
        entetes = [nettoyer_nom_colonne(c.value) or f"col{i}" for i, c in enumerate(lignes[0])]
        data, couleurs = [], []
        for row in lignes[1:]:
            vals = [c.value for c in row]
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
                continue
            data.append(vals)
            rgb = _rgb(row[0])
            couleurs.append((legende or {}).get(rgb) if rgb else None)
        df = pd.DataFrame(data, columns=entetes)
        df["_couleur"] = couleurs
        return df
    except Exception:
        raise


def lire_fsa(chemin: str) -> set[str]:
    """Lire fsa."""
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        for nom in wb.sheetnames:
            if nom.upper().startswith(C.PREFIXE_FEUILLE_FSA):
                return {
                    str(r[0]).strip().upper()
                    for r in wb[nom].iter_rows(min_row=2, max_col=1, values_only=True)
                    if r[0]
                }
        return set()
    except Exception:
        raise


def feuilles(chemin: str) -> list[str]:
    """Feuilles."""
    try:
        return load_workbook(chemin, read_only=True).sheetnames
    except Exception:
        raise


def bucket(jours: pd.Series) -> pd.Series:
    """Bucket."""
    try:
        bins = [-float("inf")] + [b[1] for b in C.BUCKETS]
        return pd.cut(jours, bins=bins, labels=C.LIBELLES_BUCKETS, right=True)
    except Exception:
        raise
